"""Deterministic, formula-aware XLSX preview parsing without executing content."""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from cra_app.edc_import.contracts import (
    EDCSheetMapping,
    EDCWorkbookMapping,
    StandardField,
)

PARSER_NAME = "openpyxl-edc-lab"
PARSER_VERSION = "1"
MAX_POPULATED_CELLS = 50000
MAX_PREVIEW_ROWS = 100000
MAX_PREVIEW_COLUMNS = 200


class ExcelImportError(Exception):
    pass


@dataclass(frozen=True)
class CellDraft:
    coordinate: str
    row: int
    column: int
    raw_text: str | None
    formula: str | None
    cached_text: str | None
    cell_type: str
    number_format: str | None

    @property
    def display_text(self) -> str | None:
        if self.formula:
            return self.cached_text
        return self.raw_text


@dataclass(frozen=True)
class SheetDraft:
    name: str
    ordinal: int
    state: str
    max_row: int
    max_column: int
    header_row: int | None
    headers: tuple[str, ...]
    populated_rows: int
    cells: tuple[CellDraft, ...]
    warnings: tuple[str, ...]


@dataclass(frozen=True)
class WorkbookDraft:
    parser_name: str
    parser_version: str
    header_signature: str
    included_sheets: tuple[str, ...]
    suggested_mapping: EDCWorkbookMapping
    sheets: tuple[SheetDraft, ...]

    @property
    def populated_cells(self) -> int:
        return sum(len(sheet.cells) for sheet in self.sheets)


_FIELD_ALIASES: dict[StandardField, set[str]] = {
    "subject_external_id": {
        "subject id",
        "subject",
        "subject number",
        "patient id",
        "patient number",
        "受试者编号",
        "患者编号",
    },
    "visit_key": {
        "visit",
        "visit name",
        "visit number",
        "访视",
        "访视编号",
    },
    "analyte": {
        "analyte",
        "test",
        "test name",
        "lab test",
        "laboratory test",
        "检验项目",
        "项目",
        "实验室项目",
    },
    "value_raw": {
        "result",
        "value",
        "lab result",
        "result value",
        "结果",
        "检验结果",
        "数值",
    },
    "unit": {"unit", "units", "单位"},
    "collected_on": {
        "collection date",
        "collected date",
        "sample date",
        "specimen date",
        "采样日期",
        "采集日期",
    },
    "collected_at": {
        "collection time",
        "collected at",
        "collection datetime",
        "sampled at",
        "采样时间",
        "采集时间",
    },
    "precision": {"precision", "date precision", "精度", "时间精度"},
    "instance_type": {
        "instance",
        "instance type",
        "test type",
        "initial or repeat",
        "initial repeat",
        "初检复查",
        "检测类型",
        "检验类型",
    },
    "repeat_instance": {
        "repeat",
        "repeat no",
        "repeat instance",
        "复查次数",
        "重复实例",
    },
    "report_on": {
        "report date",
        "reported date",
        "报告日期",
        "报告时间",
    },
    "specimen": {
        "specimen",
        "sample type",
        "specimen type",
        "样本类型",
        "标本类型",
    },
    "lower_limit": {
        "lower limit",
        "lower",
        "lln",
        "下限",
        "参考下限",
    },
    "upper_limit": {
        "upper limit",
        "upper",
        "uln",
        "上限",
        "参考上限",
    },
}


def normalize_header(value: Any) -> str:
    return re.sub(r"[\s_/\-]+", " ", str(value or "").strip().casefold())


def stringify_excel_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def _cell_type(cell, formula: str | None) -> str:
    if formula:
        return "formula"
    if isinstance(cell.value, datetime):
        return "datetime"
    if isinstance(cell.value, date):
        return "date"
    return {
        "s": "string",
        "n": "number",
        "b": "boolean",
        "e": "error",
        "d": "date",
        "f": "formula",
    }.get(cell.data_type, cell.data_type or "unknown")


def _header_values(row: tuple[Any, ...], cached_row: tuple[Any, ...]) -> list[str]:
    values = []
    for index, cell in enumerate(row):
        value = (
            cached_row[index]
            if cell.data_type == "f" and index < len(cached_row)
            else cell.value
        )
        values.append(stringify_excel_value(value) or "")
    return values


def _detect_header(
    rows: list[tuple[int, tuple[Any, ...], tuple[Any, ...]]],
    max_column: int,
) -> int | None:
    for row_number, row, cached_row in rows[:50]:
        headers = _header_values(row, cached_row)
        recognized = sum(
            1
            for value in headers
            if any(
                normalize_header(value) in aliases
                for aliases in _FIELD_ALIASES.values()
            )
        )
        nonempty = sum(bool(value.strip()) for value in headers)
        if recognized >= 2 or (nonempty >= 2 and row_number <= 5):
            return row_number
    return None


def _headers_for_row(
    rows: list[tuple[int, tuple[Any, ...], tuple[Any, ...]]],
    row_number: int | None,
    max_column: int,
) -> tuple[str, ...]:
    if row_number is None:
        return ()
    for candidate, row, cached_row in rows:
        if candidate == row_number:
            values = _header_values(row, cached_row)
            return tuple(values[:max_column])
    return ()


def suggest_sheet_mapping(
    *,
    sheet_name: str,
    header_row: int | None,
    headers: tuple[str, ...],
    included: bool,
) -> EDCSheetMapping | None:
    if header_row is None or not included:
        return None
    by_normalized = {
        normalize_header(header): header
        for header in headers
        if header.strip()
    }
    columns: dict[StandardField, str] = {}
    for field, aliases in _FIELD_ALIASES.items():
        matches = [
            by_normalized[alias]
            for alias in aliases
            if alias in by_normalized
        ]
        if matches:
            columns[field] = matches[0]
    return EDCSheetMapping(
        header_row=header_row,
        data_start_row=header_row + 1,
        columns=columns,
        included=True,
    )


def header_signature(
    sheets: list[SheetDraft], included_sheets: set[str]
) -> str:
    payload = {
        sheet.name: list(sheet.headers)
        for sheet in sheets
        if sheet.name in included_sheets
    }
    return hashlib.sha256(
        json.dumps(
            payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()


def parse_workbook(
    content: bytes,
    *,
    filename: str,
    included_sheets: list[str] | None = None,
) -> WorkbookDraft:
    if not filename.lower().endswith(".xlsx"):
        raise ExcelImportError("W07 currently supports macro-free .xlsx only")
    try:
        formula_workbook = load_workbook(
            BytesIO(content), data_only=False, read_only=True
        )
        cached_workbook = load_workbook(
            BytesIO(content), data_only=True, read_only=True
        )
    except Exception as exc:
        raise ExcelImportError(
            f"workbook could not be opened: {type(exc).__name__}"
        ) from exc

    try:
        sheets: list[SheetDraft] = []
        populated_cells = 0
        for ordinal, worksheet in enumerate(formula_workbook.worksheets, start=1):
            cached_worksheet = cached_workbook[worksheet.title]
            max_row = min(worksheet.max_row or 1, MAX_PREVIEW_ROWS)
            max_column = min(worksheet.max_column or 1, MAX_PREVIEW_COLUMNS)
            rows: list[tuple[int, tuple[Any, ...], tuple[Any, ...]]] = []
            populated_rows = 0
            for row_index, (formula_row, cached_row) in enumerate(
                zip(
                    worksheet.iter_rows(
                        min_row=1,
                        max_row=max_row,
                        max_col=max_column,
                    ),
                    cached_worksheet.iter_rows(
                        min_row=1,
                        max_row=max_row,
                        max_col=max_column,
                    ),
                ),
                start=1,
            ):
                if any(cell.value is not None for cell in formula_row):
                    populated_rows += 1
                    rows.append((row_index, formula_row, cached_row))
            header_row = _detect_header(rows, max_column)
            headers = _headers_for_row(rows, header_row, max_column)
            cells: list[CellDraft] = []
            for row_number, formula_row, cached_row in rows:
                for column_number, cell in enumerate(formula_row, start=1):
                    formula = (
                        str(cell.value)
                        if cell.data_type == "f" and cell.value is not None
                        else None
                    )
                    cached_value = (
                        cached_row[column_number - 1].value
                        if column_number <= len(cached_row)
                        else None
                    )
                    raw_text = (
                        formula
                        if formula
                        else stringify_excel_value(cell.value)
                    )
                    cached_text = (
                        stringify_excel_value(cached_value)
                        if formula
                        else None
                    )
                    if (
                        raw_text is None
                        and cached_text is None
                        and cell.value is None
                    ):
                        continue
                    cells.append(
                        CellDraft(
                            coordinate=cell.coordinate,
                            row=row_number,
                            column=column_number,
                            raw_text=raw_text,
                            formula=formula,
                            cached_text=cached_text,
                            cell_type=_cell_type(cell, formula),
                            number_format=cell.number_format,
                        )
                    )
            populated_cells += len(cells)
            if populated_cells > MAX_POPULATED_CELLS:
                raise ExcelImportError(
                    f"workbook exceeds populated-cell limit "
                    f"{MAX_POPULATED_CELLS}"
                )
            warnings = []
            if worksheet.max_row and worksheet.max_row > MAX_PREVIEW_ROWS:
                warnings.append(
                    f"sheet rows truncated at {MAX_PREVIEW_ROWS}"
                )
            if (
                worksheet.max_column
                and worksheet.max_column > MAX_PREVIEW_COLUMNS
            ):
                warnings.append(
                    f"sheet columns truncated at {MAX_PREVIEW_COLUMNS}"
                )
            if worksheet.sheet_state != "visible":
                warnings.append(f"sheet state is {worksheet.sheet_state}")
            if header_row is None:
                warnings.append("no plausible header row found")
            if any(
                cell.formula and cell.cached_text is None
                for cell in cells
            ):
                warnings.append(
                    "one or more formulas have no cached value"
                )
            sheets.append(
                SheetDraft(
                    name=worksheet.title,
                    ordinal=ordinal,
                    state=worksheet.sheet_state,
                    max_row=max_row,
                    max_column=max_column,
                    header_row=header_row,
                    headers=headers,
                    populated_rows=populated_rows,
                    cells=tuple(cells),
                    warnings=tuple(warnings),
                )
            )
    finally:
        formula_workbook.close()
        cached_workbook.close()

    known_sheets = {sheet.name for sheet in sheets}
    if included_sheets:
        missing = sorted(set(included_sheets) - known_sheets)
        if missing:
            raise ExcelImportError(
                "included sheets do not exist: " + ", ".join(missing)
            )
        included = set(included_sheets)
    else:
        included = {
            sheet.name
            for sheet in sheets
            if sheet.state == "visible" and sheet.populated_rows > 0
        }
    if not included:
        raise ExcelImportError("workbook has no included populated sheets")

    suggested: dict[str, EDCSheetMapping] = {}
    for sheet in sheets:
        mapping = suggest_sheet_mapping(
            sheet_name=sheet.name,
            header_row=sheet.header_row,
            headers=sheet.headers,
            included=sheet.name in included,
        )
        if mapping is not None:
            suggested[sheet.name] = mapping
    if not suggested:
        raise ExcelImportError(
            "included sheets have no recognizable header row"
        )
    signature = header_signature(sheets, included)
    return WorkbookDraft(
        parser_name=PARSER_NAME,
        parser_version=PARSER_VERSION,
        header_signature=signature,
        included_sheets=tuple(
            sheet.name for sheet in sheets if sheet.name in included
        ),
        suggested_mapping=EDCWorkbookMapping(sheets=suggested),
        sheets=tuple(sheets),
    )


def mapping_missing_required_fields(mapping: EDCWorkbookMapping) -> list[str]:
    missing = []
    for sheet_name, sheet in mapping.sheets.items():
        if not sheet.included:
            continue
        columns = sheet.columns
        required = {
            "subject_external_id": "subject_external_id",
            "analyte": "analyte",
            "value_raw": "value_raw",
            "collection date/time": (
                "collected_on" if "collected_on" in columns else "collected_at"
            ),
            "initial/repeat identity": (
                "instance_type"
                if "instance_type" in columns
                else "repeat_instance"
            ),
        }
        for label, field in required.items():
            if not columns.get(field):
                missing.append(f"{sheet_name}: {label}")
    return missing


def value_type_summary(sheet: SheetDraft) -> dict[str, int]:
    return dict(
        sorted(
            Counter(cell.cell_type for cell in sheet.cells).items()
        )
    )
