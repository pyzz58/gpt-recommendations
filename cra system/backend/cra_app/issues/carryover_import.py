"""Deterministic carryover extraction from report anchors, XLSX tracks and text.

The parser only proposes candidates. It does not confirm clinical meaning, invent
missing recommendations, or mutate Issues. A CRA must review and confirm the
stored preview before any Issue is created or reused.
"""
from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import load_workbook

from cra_app.issues.models import ACTION_KINDS

PARSER_NAME = "deterministic-carryover-table-text"
PARSER_VERSION = "1"


@dataclass(frozen=True)
class CandidateActionDraft:
    kind: str
    requirement: str | None = None
    action_key: str | None = None


@dataclass(frozen=True)
class CarryoverCandidateDraft:
    candidate_key: str
    ordinal: int
    original_issue_number: str | None
    original_finding: str | None
    location: str | None
    recommendation: str | None
    recommendation_sources: tuple[dict[str, Any], ...]
    known_progress: str | None
    actions: tuple[CandidateActionDraft, ...]
    source_refs: tuple[dict[str, Any], ...]
    unknown_fields: tuple[str, ...]


@dataclass(frozen=True)
class CarryoverPreview:
    parser_name: str
    parser_version: str
    candidates: tuple[CarryoverCandidateDraft, ...]
    warnings: tuple[str, ...]


_HEADER_ALIASES = {
    "issue_number": {
        "issue id",
        "issue no",
        "issue number",
        "id",
        "query id",
        "query no",
        "query number",
        "编号",
        "问题编号",
        "原编号",
        "查询编号",
    },
    "original_finding": {
        "finding",
        "description",
        "issue",
        "issue description",
        "问题",
        "发现",
        "描述",
    },
    "location": {
        "location",
        "position",
        "field",
        "site location",
        "位置",
        "表单字段",
        "条目",
    },
    "recommendation": {
        "recommendation",
        "corrective action",
        "corrective recommendation",
        "建议",
        "整改建议",
        "纠正措施",
    },
    "recommendation_sources": {
        "source",
        "sources",
        "basis",
        "recommendation source",
        "依据",
        "建议依据",
        "来源",
    },
    "known_progress": {
        "known progress",
        "progress",
        "status",
        "current status",
        "已知进展",
        "进展",
        "状态",
    },
    "action_kind": {
        "action type",
        "action kind",
        "行动类型",
        "整改类型",
    },
    "action_requirement": {
        "action requirement",
        "required action",
        "action",
        "行动要求",
        "整改要求",
        "行动",
    },
}

_LABEL_ALIASES = {
    value: key
    for key, values in _HEADER_ALIASES.items()
    for value in values
}
_KIND_ALIASES = {
    "pd": "pd_reporting",
    "pd reporting": "pd_reporting",
    "pd_reporting": "pd_reporting",
    "安全性上报": "pd_reporting",
    "edc": "edc_correction",
    "edc correction": "edc_correction",
    "edc_correction": "edc_correction",
    "edc更正": "edc_correction",
    "数据更正": "edc_correction",
    "document": "document_supply",
    "document supply": "document_supply",
    "document_supply": "document_supply",
    "补件": "document_supply",
    "medical clarification": "medical_clarification",
    "medical_clarification": "medical_clarification",
    "医学澄清": "medical_clarification",
    "query": "query_followup",
    "query followup": "query_followup",
    "query_followup": "query_followup",
    "other": "other",
    "其他": "other",
}


def _normalize_header(value: Any) -> str:
    return re.sub(r"[\s_/\-]+", " ", str(value or "").strip().casefold())


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        cleaned = value.isoformat()
    else:
        cleaned = str(value).strip()
    return cleaned or None


def _split_actions(kind_value: str | None, requirement: str | None) -> list[CandidateActionDraft]:
    if not kind_value and not requirement:
        return []
    kinds = [
        item.strip()
        for item in re.split(r"[,，;；|]", kind_value or "")
        if item.strip()
    ]
    if not kinds:
        kinds = ["other"]
    result = []
    requirements = [
        item.strip()
        for item in re.split(r"\s*[|；;]\s*", requirement or "")
        if item.strip()
    ]
    for index, raw_kind in enumerate(kinds):
        kind = _KIND_ALIASES.get(raw_kind.casefold(), raw_kind.casefold())
        if kind not in ACTION_KINDS:
            kind = "other"
        result.append(
            CandidateActionDraft(
                kind=kind,
                requirement=(
                    requirements[index]
                    if index < len(requirements)
                    else requirements[0]
                    if len(kinds) == 1 and requirements
                    else None
                ),
                action_key=f"A{index + 1}",
            )
        )
    return result


def _source_ref(anchor: Any, *, label: str | None = None) -> dict[str, Any]:
    return {
        "source_type": "document_anchor",
        "document_version_id": str(
            getattr(anchor, "document_version_id", "")
        ),
        "anchor_id": str(getattr(anchor, "id", "")),
        "anchor_type": getattr(anchor, "anchor_type", None),
        "locator": getattr(anchor, "logical_node_id", None),
        "page_number": getattr(anchor, "page_number", None),
        "table_index": getattr(anchor, "table_index", None),
        "row_index": getattr(anchor, "row_index", None),
        "column_index": getattr(anchor, "column_index", None),
        "label": label,
        "quote": (getattr(anchor, "text", "") or "")[:1000],
    }


def _candidate_key(
    source_sha256: str,
    ordinal: int,
    issue_number: str | None,
    locator: str,
) -> str:
    payload = (
        f"{source_sha256}|{ordinal}|{issue_number or ''}|{locator}".encode(
            "utf-8"
        )
    )
    return hashlib.sha256(payload).hexdigest()


def _draft(
    *,
    source_sha256: str,
    ordinal: int,
    locator: str,
    values: dict[str, str | None],
    source_refs: list[dict[str, Any]],
) -> CarryoverCandidateDraft | None:
    issue_number = _clean_text(values.get("issue_number"))
    finding = _clean_text(values.get("original_finding"))
    if not issue_number and not finding:
        return None
    location = _clean_text(values.get("location"))
    recommendation = _clean_text(values.get("recommendation"))
    known_progress = _clean_text(values.get("known_progress"))
    sources_text = _clean_text(values.get("recommendation_sources"))
    recommendation_sources: tuple[dict[str, Any], ...] = ()
    if sources_text:
        recommendation_sources = (
            {
                "text": sources_text,
                "source_refs": source_refs,
                "extraction": "source_text",
            },
        )
    actions = tuple(
        _split_actions(values.get("action_kind"), values.get("action_requirement"))
    )
    unknown_fields = tuple(
        field_name
        for field_name, value in (
            ("original_issue_number", issue_number),
            ("original_finding", finding),
            ("location", location),
            ("recommendation", recommendation),
            ("known_progress", known_progress),
        )
        if not value
    )
    if not actions:
        unknown_fields += ("actions",)
    return CarryoverCandidateDraft(
        candidate_key=_candidate_key(
            source_sha256, ordinal, issue_number, locator
        ),
        ordinal=ordinal,
        original_issue_number=issue_number,
        original_finding=finding,
        location=location,
        recommendation=recommendation,
        recommendation_sources=recommendation_sources,
        known_progress=known_progress,
        actions=actions,
        source_refs=tuple(source_refs),
        unknown_fields=unknown_fields,
    )


def _table_candidates(
    anchors: Iterable[Any], source_sha256: str
) -> list[CarryoverCandidateDraft]:
    grouped: dict[tuple[int, int], list[Any]] = defaultdict(list)
    for anchor in anchors:
        if anchor.table_index is None or anchor.row_index is None:
            continue
        grouped[(anchor.table_index, anchor.row_index)].append(anchor)
    candidates: list[CarryoverCandidateDraft] = []
    by_table: dict[int, list[tuple[int, list[Any]]]] = defaultdict(list)
    for (table_index, row_index), cells in grouped.items():
        by_table[table_index].append((row_index, cells))
    for table_index in sorted(by_table):
        rows = sorted(by_table[table_index], key=lambda item: item[0])
        header: dict[int, str] | None = None
        mapped_headers: dict[str, int] = {}
        for row_index, cells in rows:
            row_mapping: dict[int, str] = {}
            for cell in sorted(cells, key=lambda item: item.column_index or 0):
                normalized = _normalize_header(cell.text)
                field_name = _LABEL_ALIASES.get(normalized)
                if field_name:
                    row_mapping[cell.column_index or 0] = field_name
            if len(set(row_mapping.values())) >= 2:
                header = row_mapping
                mapped_headers = {
                    field_name: column
                    for column, field_name in row_mapping.items()
                }
                break
        if header is None:
            continue
        header_row_index = next(
            row_index
            for row_index, cells in rows
            if any(
                _LABEL_ALIASES.get(_normalize_header(cell.text))
                for cell in cells
            )
        )
        for ordinal, (row_index, cells) in enumerate(rows):
            if row_index == header_row_index:
                continue
            values: dict[str, str | None] = {}
            refs = []
            for cell in sorted(cells, key=lambda item: item.column_index or 0):
                column = cell.column_index or 0
                field_name = header.get(column)
                if field_name:
                    values[field_name] = _clean_text(cell.text)
                refs.append(_source_ref(cell))
            candidate = _draft(
                source_sha256=source_sha256,
                ordinal=ordinal,
                locator=f"table[{table_index}]/row[{row_index}]",
                values=values,
                source_refs=refs,
            )
            if candidate is not None:
                candidates.append(candidate)
    return candidates


def _block_candidates(
    anchors: Iterable[Any], source_sha256: str
) -> list[CarryoverCandidateDraft]:
    blocks: list[tuple[str, list[Any]]] = []
    current_refs: list[Any] = []
    current_locator = ""
    current_values: dict[str, str | None] = {}

    def flush() -> None:
        nonlocal current_refs, current_locator, current_values
        if current_refs:
            blocks.append(
                (
                    current_locator,
                    [
                        *current_refs,
                    ],
                )
            )
            current_values = {}
        current_refs = []
        current_locator = ""

    for anchor in anchors:
        if anchor.table_index is not None:
            continue
        text = (anchor.text or "").strip()
        if not text:
            continue
        label = None
        value = text
        for separator in (":", "："):
            if separator in text:
                possible, remainder = text.split(separator, 1)
                normalized = _normalize_header(possible)
                if normalized in _LABEL_ALIASES:
                    label = _LABEL_ALIASES[normalized]
                    value = remainder.strip()
                    break
        issue_match = re.search(
            r"\b(?:ISSUE(?:\s+ID)?|QUERY)\s*[:：#-]?\s*([A-Za-z0-9_.-]+)",
            text,
            flags=re.IGNORECASE,
        )
        if issue_match and label != "original_finding":
            flush()
            current_locator = getattr(anchor, "logical_node_id", "")
            current_values["issue_number"] = issue_match.group(1)
        elif label == "issue_number":
            flush()
            current_locator = getattr(anchor, "logical_node_id", "")
            current_values["issue_number"] = value
        if not current_refs:
            current_locator = (
                current_locator or getattr(anchor, "logical_node_id", "")
            )
        if label and label != "issue_number":
            current_values[label] = value
        current_refs.append(anchor)
    flush()

    candidates = []
    for ordinal, (locator, refs) in enumerate(blocks):
        values: dict[str, str | None] = {}
        for anchor in refs:
            text = (anchor.text or "").strip()
            for separator in (":", "："):
                if separator not in text:
                    continue
                possible, remainder = text.split(separator, 1)
                normalized = _normalize_header(possible)
                field_name = _LABEL_ALIASES.get(normalized)
                if field_name:
                    values[field_name] = remainder.strip()
                    break
            issue_match = re.search(
                r"\b(?:ISSUE(?:\s+ID)?|QUERY)\s*[:：#-]?\s*([A-Za-z0-9_.-]+)",
                text,
                flags=re.IGNORECASE,
            )
            if issue_match and "issue_number" not in values:
                values["issue_number"] = issue_match.group(1)
        candidate = _draft(
            source_sha256=source_sha256,
            ordinal=ordinal,
            locator=locator,
            values=values,
            source_refs=[_source_ref(anchor) for anchor in refs],
        )
        if candidate is not None:
            candidates.append(candidate)
    return candidates


def preview_anchors(
    *,
    anchors: list[Any],
    source_sha256: str,
) -> CarryoverPreview:
    table_candidates = _table_candidates(anchors, source_sha256)
    block_candidates = _block_candidates(anchors, source_sha256)
    merged = list(table_candidates)
    for candidate in block_candidates:
        existing_index = next(
            (
                index
                for index, item in enumerate(merged)
                if candidate.original_issue_number
                and item.original_issue_number
                == candidate.original_issue_number
            ),
            None,
        )
        if existing_index is None:
            merged.append(candidate)
            continue
        existing = merged[existing_index]
        source_refs = {
            (
                ref.get("document_version_id"),
                ref.get("anchor_id"),
                ref.get("locator"),
                ref.get("quote"),
            ): ref
            for ref in [*existing.source_refs, *candidate.source_refs]
        }
        actions = {
            (action.kind, action.requirement): action
            for action in [*existing.actions, *candidate.actions]
        }
        sources = {
            json.dumps(item, ensure_ascii=False, sort_keys=True): item
            for item in [
                *existing.recommendation_sources,
                *candidate.recommendation_sources,
            ]
        }
        merged[existing_index] = replace(
            existing,
            original_finding=(
                existing.original_finding or candidate.original_finding
            ),
            location=existing.location or candidate.location,
            recommendation=existing.recommendation or candidate.recommendation,
            recommendation_sources=tuple(sources.values()),
            known_progress=(
                existing.known_progress or candidate.known_progress
            ),
            actions=tuple(actions.values()),
            source_refs=tuple(source_refs.values()),
            unknown_fields=tuple(
                field_name
                for field_name, value in (
                    (
                        "original_issue_number",
                        existing.original_issue_number
                        or candidate.original_issue_number,
                    ),
                    (
                        "original_finding",
                        existing.original_finding
                        or candidate.original_finding,
                    ),
                    ("location", existing.location or candidate.location),
                    (
                        "recommendation",
                        existing.recommendation or candidate.recommendation,
                    ),
                    (
                        "known_progress",
                        existing.known_progress or candidate.known_progress,
                    ),
                )
                if not value
            )
            + (() if actions else ("actions",)),
        )
    candidates = merged
    warnings: list[str] = []
    if not candidates:
        warnings.append("no_carryover_candidates")
    if table_candidates and block_candidates:
        warnings.append("table_and_paragraph_candidates_merged")
    return CarryoverPreview(
        parser_name=PARSER_NAME,
        parser_version=PARSER_VERSION,
        candidates=tuple(candidates),
        warnings=tuple(warnings),
    )


def _mapped_xlsx_header(values: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, value in enumerate(values):
        field_name = _LABEL_ALIASES.get(_normalize_header(value))
        if field_name:
            mapping[index] = field_name
    return mapping


def preview_xlsx(
    content: bytes,
    *,
    source_sha256: str,
    filename: str,
) -> CarryoverPreview:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    candidates: list[CarryoverCandidateDraft] = []
    warnings: list[str] = []
    ordinal = 0
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            header_index = None
            header: dict[int, str] = {}
            for index, row in enumerate(rows):
                mapped = _mapped_xlsx_header(list(row))
                if len(set(mapped.values())) >= 2:
                    header_index = index
                    header = mapped
                    break
            if header_index is None:
                warnings.append(f"{worksheet.title}: no recognized issue columns")
                continue
            for row_number, row in enumerate(
                rows[header_index + 1 :], start=header_index + 2
            ):
                values: dict[str, str | None] = {}
                refs = []
                for column_index, value in enumerate(row):
                    field_name = header.get(column_index)
                    if field_name:
                        values[field_name] = _clean_text(value)
                    if value is not None:
                        refs.append(
                            {
                                "source_type": "xlsx_cell",
                                "filename": filename,
                                "sheet": worksheet.title,
                                "row": row_number,
                                "column": column_index + 1,
                                "label": field_name,
                                "quote": str(value)[:1000],
                            }
                        )
                candidate = _draft(
                    source_sha256=source_sha256,
                    ordinal=ordinal,
                    locator=f"{worksheet.title}!row[{row_number}]",
                    values=values,
                    source_refs=refs,
                )
                if candidate is not None:
                    candidates.append(candidate)
                    ordinal += 1
    finally:
        workbook.close()
    if not candidates:
        warnings.append("no_carryover_candidates")
    return CarryoverPreview(
        parser_name=f"{PARSER_NAME}-xlsx",
        parser_version=PARSER_VERSION,
        candidates=tuple(candidates),
        warnings=tuple(warnings),
    )


def preview_text(
    text: str,
    *,
    source_sha256: str,
    source_label: str,
) -> CarryoverPreview:
    lines = text.splitlines()
    anchors = [
        type(
            "TextAnchor",
            (),
            {
                "id": "",
                "document_version_id": "",
                "anchor_type": "plain_text",
                "logical_node_id": f"{source_label}:line[{index}]",
                "text": line,
                "page_number": None,
                "table_index": None,
                "row_index": None,
                "column_index": None,
            },
        )()
        for index, line in enumerate(lines, start=1)
        if line.strip()
    ]
    preview = preview_anchors(anchors=anchors, source_sha256=source_sha256)
    return CarryoverPreview(
        parser_name=f"{PARSER_NAME}-text",
        parser_version=PARSER_VERSION,
        candidates=preview.candidates,
        warnings=preview.warnings,
    )
