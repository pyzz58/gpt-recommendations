"""Parse authorized structured JSON, CSV or XLSX context imports."""
from __future__ import annotations

import csv
import json
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from cra_app.subjects.contracts import FactImportRow

MAX_CONTEXT_IMPORT_BYTES = 5 * 1024 * 1024
REQUIRED_COLUMNS = {
    "subject_external_id",
    "fact_key",
    "value_raw",
    "observed_on",
    "source_ref",
}
OPTIONAL_COLUMNS = {
    "value_type",
    "precision",
    "unit",
    "repeat_instance",
    "visit_key",
    "source_kind",
    "verification_status",
}


class ContextImportError(ValueError):
    pass


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_source_ref(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if not value:
        raise ContextImportError("source_ref is required")
    try:
        parsed = json.loads(str(value))
    except json.JSONDecodeError as exc:
        raise ContextImportError("source_ref must be JSON") from exc
    if not isinstance(parsed, dict):
        raise ContextImportError("source_ref must be a JSON object")
    return parsed


def _normalize_mapping(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {}
    for key, value in row.items():
        if key is None:
            continue
        if key == "source_ref" and isinstance(value, dict):
            normalized[key] = value
        elif _cell_text(value):
            normalized[key] = _cell_text(value)
    if "source_ref" in normalized:
        normalized["source_ref"] = _parse_source_ref(normalized["source_ref"])
    return normalized


def rows_from_json(content: bytes) -> list[FactImportRow]:
    try:
        value = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ContextImportError("invalid JSON context import") from exc
    rows = value.get("rows") if isinstance(value, dict) else value
    if not isinstance(rows, list) or not rows:
        raise ContextImportError("JSON context import requires a non-empty rows list")
    return [FactImportRow.model_validate(_normalize_mapping(row)) for row in rows]


def rows_from_csv(content: bytes) -> list[FactImportRow]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ContextImportError("CSV context import must be UTF-8") from exc
    reader = csv.DictReader(StringIO(text))
    fields = set(reader.fieldnames or ())
    missing = REQUIRED_COLUMNS - fields
    if missing:
        raise ContextImportError(
            "CSV context import missing columns: " + ", ".join(sorted(missing))
        )
    rows = [_normalize_mapping(row) for row in reader]
    if not rows:
        raise ContextImportError("CSV context import contains no facts")
    return [FactImportRow.model_validate(row) for row in rows]


def rows_from_xlsx(content: bytes) -> list[FactImportRow]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ContextImportError("invalid XLSX context import") from exc
    worksheet = workbook.active
    values = worksheet.iter_rows(values_only=True)
    try:
        headers = [_cell_text(value) for value in next(values)]
    except StopIteration as exc:
        raise ContextImportError("XLSX context import has no header") from exc
    field_names = set(headers)
    missing = REQUIRED_COLUMNS - field_names
    if missing:
        raise ContextImportError(
            "XLSX context import missing columns: " + ", ".join(sorted(missing))
        )
    rows = []
    for values_row in values:
        values_map = {
            header: value
            for header, value in zip(headers, values_row, strict=False)
            if header
        }
        if not any(_cell_text(value) for value in values_map.values()):
            continue
        rows.append(FactImportRow.model_validate(_normalize_mapping(values_map)))
    if not rows:
        raise ContextImportError("XLSX context import contains no facts")
    return rows


def parse_context_file(filename: str, content: bytes) -> list[FactImportRow]:
    if not content:
        raise ContextImportError("context import file is empty")
    if len(content) > MAX_CONTEXT_IMPORT_BYTES:
        raise ContextImportError("context import file exceeds 5 MiB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        return rows_from_json(content)
    if suffix == ".csv":
        return rows_from_csv(content)
    if suffix == ".xlsx":
        return rows_from_xlsx(content)
    raise ContextImportError("context import supports JSON, CSV and XLSX")
