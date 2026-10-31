"""Generate the public, synthetic W03 source corpus deterministically enough.

The generated files are committed. Tests never download clinical material or
call a provider while creating fixtures.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artifacts"


def _synthetic_docx(path: Path) -> None:
    document = Document()
    document.core_properties.title = "Synthetic protocol excerpt"
    document.core_properties.author = "Synthetic fixture generator"
    document.core_properties.created = datetime(2026, 10, 31)
    document.core_properties.modified = datetime(2026, 10, 31)
    document.add_heading("Synthetic Study Protocol", level=0)
    document.add_heading("1. Screening", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Consent must be obtained before ")
    paragraph.add_run("any study procedure").bold = True
    paragraph.add_run(" begins. This sentence is synthetic.")
    document.add_heading("2. Laboratory", level=1)
    run = document.add_paragraph().add_run(
        "If local laboratory reference ranges are used, verify the applicable "
        "project rule. "
    )
    run.italic = True
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Visit"
    table.cell(0, 1).text = "Window"
    table.cell(0, 2).text = "Condition"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "V1"
    table.cell(1, 1).text = "±3 days"
    table.cell(1, 2).text = "Synthetic condition C -> path B"
    table.cell(2, 0).text = "V2"
    table.cell(2, 1).text = "±5 days"
    table.cell(2, 2).text = "Synthetic condition D -> path F"
    media = Image.new("RGB", (24, 12), "white")
    media_bytes = BytesIO()
    media.save(media_bytes, format="PNG")
    media_bytes.seek(0)
    document.add_picture(media_bytes)
    document.add_heading("3. Background", level=1)
    document.add_paragraph(
        "This is explicitly marked background content and is not a requirement."
    )
    document.add_paragraph("Synthetic footnote target.")
    document.save(path)

    with tempfile.TemporaryDirectory() as temporary:
        temp = Path(temporary)
        with zipfile.ZipFile(path) as source:
            source.extractall(temp)
        document_xml = temp / "word/document.xml"
        xml = document_xml.read_text(encoding="utf-8")
        target = (
            "<w:p><w:r><w:t>This is explicitly marked background content "
            "and is not a requirement.</w:t></w:r></w:p>"
        )
        if target not in xml:
            raise RuntimeError("background fixture paragraph not found")
        wrapped = (
            '<w:sdt><w:sdtPr><w:tag w:val="CRA-BACKGROUND"/>'
            "<w:alias w:val=\"Background\"/></w:sdtPr><w:sdtContent>"
            + target
            + "</w:sdtContent></w:sdt>"
        )
        xml = xml.replace(target, wrapped, 1)
        footnote_target = (
            "<w:p><w:r><w:t>Synthetic footnote target.</w:t></w:r></w:p>"
        )
        if footnote_target not in xml:
            raise RuntimeError("footnote fixture paragraph not found")
        xml = xml.replace(
            footnote_target,
            (
                "<w:p><w:r><w:t>Synthetic footnote target.</w:t></w:r>"
                '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
                '<w:footnoteReference w:id="1"/></w:r></w:p>'
            ),
            1,
        )
        document_xml.write_text(xml, encoding="utf-8")
        (temp / "word/footnotes.xml").write_text(
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>Synthetic footnote: applicability is illustrative only.</w:t></w:r></w:p></w:footnote>
</w:footnotes>
""",
            encoding="utf-8",
        )
        content_types = temp / "[Content_Types].xml"
        content_xml = content_types.read_text(encoding="utf-8")
        if "footnotes+xml" not in content_xml:
            content_xml = content_xml.replace(
                "</Types>",
                '<Override PartName="/word/footnotes.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.'
                'wordprocessingml.footnotes+xml"/></Types>',
            )
            content_types.write_text(content_xml, encoding="utf-8")
        rels = temp / "word/_rels/document.xml.rels"
        rels_xml = rels.read_text(encoding="utf-8")
        if "footnotes.xml" not in rels_xml:
            rels_xml = rels_xml.replace(
                "</Relationships>",
                '<Relationship Id="rIdCRAFootnotes" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/footnotes" Target="footnotes.xml"/></Relationships>',
            )
            rels.write_text(rels_xml, encoding="utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in sorted(temp.rglob("*")):
                if item.is_file():
                    archive.write(item, item.relative_to(temp).as_posix())


def _synthetic_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("Synthetic monitoring report")
    pdf.setAuthor("Synthetic fixture generator")
    width, height = A4
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(60, height - 70, "Synthetic Monitoring Report")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(60, height - 95, "Site: Synthetic Site A    Subject: SYN-001")
    pdf.drawString(60, height - 115, "This file contains no real clinical information.")
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(60, height - 155, "Laboratory review")
    text = "Verify the locally reported value against the synthetic EDC export."
    pdf.setFont("Helvetica", 11)
    pdf.drawString(60, height - 175, text)
    left, top, right, bottom = 60, height - 210, width - 60, height - 330
    widths = [100, 100, right - left - 200]
    heights = [30] * 4
    pdf.setLineWidth(0.8)
    y = top
    for row_height in heights:
        pdf.line(left, y, right, y)
        y -= row_height
    pdf.line(left, bottom, right, bottom)
    x = left
    for cell_width in widths:
        pdf.line(x, top, x, bottom)
        x += cell_width
    pdf.line(right, top, right, bottom)
    rows = [
        ("Analyte", "Result", "Comment"),
        ("Synthetic Hb", "121 g/L", "Within synthetic local range"),
        ("Synthetic ALT", "31 U/L", "Synthetic only"),
    ]
    for row_index, row in enumerate(rows):
        y = top - 20 - row_index * 30
        x = left + 8
        for column_index, value in enumerate(row):
            font = "Helvetica-Bold" if row_index == 0 else "Helvetica"
            pdf.setFont(font, 9)
            pdf.drawString(x, y, value)
            x += widths[column_index]
    pdf.showPage()
    pdf.save()


def _synthetic_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Carries"
    sheet.append(["Issue ID", "Description", "Status"])
    sheet.append(["CRA-SYN-001", "Synthetic carryover example", "Open"])
    sheet.append(["CRA-SYN-002", "Synthetic directory example", "Pending"])
    workbook.properties.creator = "Synthetic fixture generator"
    workbook.save(path)


def _synthetic_image(path: Path) -> None:
    image = Image.new("RGB", (1200, 600), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 1170, 570), outline="black", width=3)
    draw.text((70, 70), "SYNTHETIC LAB REPORT", fill="black")
    draw.text((70, 140), "Analyte      Result      Unit", fill="black")
    draw.text((70, 210), "SYNTH-Hb     121         g/L", fill="black")
    draw.text((70, 500), "NO REAL PATIENT DATA", fill="black")
    image.save(path)


def _hashes() -> dict[str, str]:
    return {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(OUTPUT.iterdir())
        if path.is_file()
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for path in OUTPUT.glob("*"):
        if path.is_file():
            path.unlink()
    protocol = OUTPUT / "synthetic_rule_protocol.docx"
    _synthetic_docx(protocol)
    shutil.copy2(protocol, OUTPUT / "synthetic_report_template.docx")
    _synthetic_pdf(OUTPUT / "synthetic_monitoring_report.pdf")
    _synthetic_xlsx(OUTPUT / "synthetic_track.xlsx")
    _synthetic_image(OUTPUT / "synthetic_lab_report.png")
    (OUTPUT / "synthetic_rule_text.txt").write_text(
        "Synthetic rule text\n"
        "Background: this line is synthetic context.\n"
        "Requirement: if condition C is present, follow path B.\n",
        encoding="utf-8",
    )
    hashes = _hashes()
    manifest = {
        "schema_version": 1,
        "source": "pure_synthetic",
        "generated_on": str(date(2026, 10, 31)),
        "generator": "tests/fixtures/documents/generate_fixtures.py",
        "artifacts": [
            {
                "artifact_id": "RULE-DOCX-DEV",
                "file": "artifacts/synthetic_rule_protocol.docx",
                "format": "docx",
                "purpose": "rule",
                "group": "development",
                "sha256": hashes["synthetic_rule_protocol.docx"],
                "expected": {
                    "anchors": [
                        "table_cell",
                        "footnote",
                        "paragraph",
                        "media",
                    ],
                    "section_path": ["Synthetic Study Protocol", "1. Screening"],
                    "background_tag": "CRA-BACKGROUND",
                    "bold_text": "any study procedure",
                    "footnote_text": "Synthetic footnote: applicability is illustrative only.",
                },
            },
            {
                "artifact_id": "REPORT-PDF-EVAL",
                "file": "artifacts/synthetic_monitoring_report.pdf",
                "format": "pdf",
                "purpose": "carryover_report",
                "group": "evaluation",
                "sha256": hashes["synthetic_monitoring_report.pdf"],
                "expected": {
                    "page": 1,
                    "anchor_types": ["pdf_text", "table_cell"],
                    "contains": ["SYN-001", "Synthetic Hb"],
                    "coordinates_present": True,
                },
            },
            {
                "artifact_id": "TRACK-XLSX-DEV",
                "file": "artifacts/synthetic_track.xlsx",
                "format": "xlsx",
                "purpose": "track",
                "group": "development",
                "sha256": hashes["synthetic_track.xlsx"],
                "expected": {
                    "retained_as_private_file": True,
                    "structured_parsing_task": "CRA-W07/W06",
                },
            },
            {
                "artifact_id": "LAB-PNG-PROBE",
                "file": "artifacts/synthetic_lab_report.png",
                "format": "png",
                "purpose": "fact",
                "group": "development",
                "sha256": hashes["synthetic_lab_report.png"],
                "expected": {
                    "ocr_candidate_probe_only": True,
                    "minimum_expected_tokens": ["SYNTH-Hb", "121", "g/L"],
                },
            },
            {
                "artifact_id": "TEMPLATE-DOCX-DEV",
                "file": "artifacts/synthetic_report_template.docx",
                "format": "docx",
                "purpose": "template",
                "group": "development",
                "sha256": hashes["synthetic_report_template.docx"],
                "expected": {"retained_as_private_file": True},
            },
            {
                "artifact_id": "RULE-TEXT-DEV",
                "file": "artifacts/synthetic_rule_text.txt",
                "format": "txt",
                "purpose": "rule",
                "group": "development",
                "sha256": hashes["synthetic_rule_text.txt"],
                "expected": {
                    "anchor_types": ["plain_text"],
                    "contains": ["condition C", "path B"],
                },
            },
        ],
    }
    (ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
