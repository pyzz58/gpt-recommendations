"""Generate a public synthetic multi-sheet EDC laboratory export."""
from __future__ import annotations

import hashlib
import json
import re
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artifacts"
MANIFEST = ROOT / "manifest.json"

HEADERS = [
    "Subject ID",
    "Visit",
    "Analyte",
    "Result",
    "Unit",
    "Collection Date",
    "Collection Time",
    "Instance",
    "Repeat No",
    "Report Date",
    "Specimen",
    "LLN",
    "ULN",
]


def _append(
    sheet,
    subject: str,
    visit: str,
    analyte: str,
    result,
    unit: str | None,
    instance: str,
    repeat: str,
) -> None:
    sheet.append(
        [
            subject,
            visit,
            analyte,
            result,
            unit,
            date(2026, 1, 2),
            time(9, 15 if instance == "Initial" else 45),
            instance,
            repeat,
            date(2026, 1, 2),
            "Blood",
            "110",
            "160",
        ]
    )


def _workbook(path: Path) -> None:
    workbook = Workbook()
    hematology = workbook.active
    hematology.title = "Hematology"
    hematology.append(HEADERS)
    _append(
        hematology,
        "SYN-EDC-001",
        "V1",
        "Hemoglobin",
        121,
        "g/L",
        "Initial",
        "0",
    )
    _append(
        hematology,
        "SYN-EDC-001",
        "V1",
        "Hemoglobin",
        "=IF(1=1,119,0)",
        "g/L",
        "Repeat",
        "1",
    )
    _append(
        hematology,
        "SYN-EDC-002",
        "V2",
        "Creatinine",
        88,
        "umol/L",
        "Initial",
        "0",
    )

    chemistry = workbook.create_sheet("Chemistry")
    chemistry.append(HEADERS)
    _append(
        chemistry,
        "SYN-EDC-002",
        "V2",
        "ALT",
        "=16*2",
        "U/L",
        "Initial",
        "0",
    )
    _append(
        chemistry,
        "SYN-EDC-001",
        "V1",
        "Glucose",
        5.1,
        None,
        "Initial",
        "0",
    )

    notes = workbook.create_sheet("Internal Notes")
    notes.sheet_state = "hidden"
    notes.append(["Note"])
    notes.append(["This hidden sheet is not part of the laboratory export."])
    workbook.properties.creator = "Synthetic fixture generator"
    workbook.properties.created = datetime(2026, 10, 31)
    workbook.properties.modified = datetime(2026, 10, 31)
    workbook.save(path)


def _inject_cached_formula(path: Path) -> None:
    with tempfile.TemporaryDirectory() as temporary:
        temp = Path(temporary)
        with zipfile.ZipFile(path) as source:
            source.extractall(temp)
        sheet = temp / "xl/worksheets/sheet1.xml"
        xml = sheet.read_text(encoding="utf-8")
        pattern = re.compile(
            r'(<c r="D3"[^>]*><f>IF\(1=1,119,0\)</f><v>)\s*(</v>)'
        )
        xml, replacements = pattern.subn(r"\g<1>119\g<2>", xml, count=1)
        if replacements != 1:
            raise RuntimeError("formula cache target was not found")
        sheet.write_text(xml, encoding="utf-8")
        with zipfile.ZipFile(
            path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for item in sorted(temp.rglob("*")):
                if not item.is_file():
                    continue
                name = item.relative_to(temp).as_posix()
                info = zipfile.ZipInfo(name, date_time=(2026, 10, 31, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, item.read_bytes())


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    artifact = OUTPUT / "synthetic_edc_lab.xlsx"
    _workbook(artifact)
    _inject_cached_formula(artifact)
    content = artifact.read_bytes()
    MANIFEST.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "source": "pure_synthetic",
                "generator": "tests/fixtures/edc/generate_fixtures.py",
                "artifacts": [
                    {
                        "artifact_id": "EDC-LAB-XLSX-DEV",
                        "file": "edc/artifacts/synthetic_edc_lab.xlsx",
                        "format": "xlsx",
                        "sha256": hashlib.sha256(content).hexdigest(),
                        "expected": {
                            "sheets": [
                                "Hematology",
                                "Chemistry",
                                "Internal Notes",
                            ],
                            "cached_formula": {
                                "coordinate": "Hematology!D3",
                                "formula": "=IF(1=1,119,0)",
                                "cached": "119",
                            },
                            "missing_formula_cache": "Chemistry!D2",
                            "same_day_instances": ["Initial", "Repeat"],
                        },
                    }
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
