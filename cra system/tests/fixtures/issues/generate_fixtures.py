"""Generate public synthetic carryover report and track fixtures."""
from __future__ import annotations

import hashlib
import json
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artifacts"
MANIFEST = ROOT / "manifest.json"


def _report(path: Path) -> None:
    document = Document()
    document.core_properties.title = "Synthetic carryover monitoring report"
    document.core_properties.author = "Synthetic fixture generator"
    document.core_properties.created = datetime(2026, 10, 31)
    document.core_properties.modified = datetime(2026, 10, 31)
    document.add_heading("Synthetic Previous Monitoring Report", level=0)
    document.add_heading("Outstanding findings", level=1)

    document.add_paragraph("Issue ID: CRA-SYN-101")
    document.add_paragraph(
        "Finding: Synthetic consent date was inconsistent with the source."
    )
    document.add_paragraph("Location: Consent form, signature date")
    document.add_paragraph(
        "Recommendation: Correct the synthetic date and retain source evidence."
    )
    document.add_paragraph("Recommendation Source: Synthetic SOP section 4.1")
    document.add_paragraph(
        "Known Progress: Site reported a correction; CRA review remains open."
    )
    document.add_paragraph("Action Type: edc_correction")
    document.add_paragraph(
        "Action Requirement: Verify the corrected synthetic consent date."
    )

    document.add_paragraph("Issue ID: CRA-SYN-102")
    document.add_paragraph(
        "Finding: A synthetic serious event was reported outside the example "
        "reporting window."
    )
    document.add_paragraph("Known Progress: Reporting status not stated.")
    document.add_paragraph("Action Type: pd_reporting")
    document.add_paragraph(
        "Action Requirement: Confirm whether the synthetic PD report was sent."
    )
    document.save(path)


def _track(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Carryover"
    sheet.append(
        [
            "Issue ID",
            "Finding",
            "Location",
            "Recommendation",
            "Recommendation Source",
            "Known Progress",
            "Action Type",
            "Action Requirement",
        ]
    )
    sheet.append(
        [
            "CRA-SYN-101",
            "Synthetic consent date was inconsistent with the source.",
            "Consent form, signature date",
            "Correct the synthetic date and retain source evidence.",
            "Synthetic SOP section 4.1",
            "Site reported correction; follow-up review due.",
            "edc_correction",
            "Verify the corrected synthetic consent date.",
        ]
    )
    sheet.append(
        [
            "CRA-SYN-103",
            "Synthetic training record was not available in the supplied folder.",
            None,
            None,
            None,
            "Folder state unknown.",
            "document_supply",
            "Obtain the synthetic training record if applicable.",
        ]
    )
    workbook.properties.creator = "Synthetic fixture generator"
    workbook.save(path)


def _normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory() as temporary:
        normalized = Path(temporary) / path.name
        with zipfile.ZipFile(path) as source:
            entries = [
                (
                    info.filename,
                    source.read(info.filename),
                    info.compress_type,
                )
                for info in source.infolist()
            ]
        with zipfile.ZipFile(
            normalized, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for name, content, compression in sorted(entries):
                info = zipfile.ZipInfo(name, date_time=(2026, 10, 31, 0, 0, 0))
                info.compress_type = compression
                info.external_attr = 0o600 << 16
                target.writestr(info, content)
        normalized.replace(path)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report = OUTPUT / "synthetic_carryover_report.docx"
    track = OUTPUT / "synthetic_carryover_track.xlsx"
    _report(report)
    _track(track)
    _normalize_ooxml(report)
    _normalize_ooxml(track)
    artifacts = []
    for artifact_id, path, purpose, expected in (
        (
            "CARRYOVER-DOCX-DEV",
            report,
            "carryover_report",
            {"issue_numbers": ["CRA-SYN-101", "CRA-SYN-102"]},
        ),
        (
            "CARRYOVER-XLSX-DEV",
            track,
            "track",
            {
                "issue_numbers": ["CRA-SYN-101", "CRA-SYN-103"],
                "unknown_fields": ["location", "recommendation"],
            },
        ),
    ):
        content = path.read_bytes()
        artifacts.append(
            {
                "artifact_id": artifact_id,
                "file": f"issues/artifacts/{path.name}",
                "purpose": purpose,
                "format": path.suffix[1:],
                "sha256": hashlib.sha256(content).hexdigest(),
                "expected": expected,
            }
        )
    MANIFEST.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "source": "pure_synthetic",
                "generator": "tests/fixtures/issues/generate_fixtures.py",
                "artifacts": artifacts,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
